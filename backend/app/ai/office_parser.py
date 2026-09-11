from __future__ import annotations

import re
from io import BytesIO
from typing import List

from app.ai.types import PageText, TextBlock
from app.core.exceptions import ValidationAppError


def extract_docx_pages(data: bytes) -> List[PageText]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValidationAppError("DOCX support is not installed on this server.") from exc
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ValidationAppError("This Word document could not be read.") from exc

    heading = None
    blocks: List[TextBlock] = []
    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        kind = "paragraph"
        if "Heading" in style or "Title" in style:
            kind = "heading"
            heading = text[:180]
        elif paragraph.style and "List" in (paragraph.style.name or ""):
            kind = "list"
        blocks.append(TextBlock(text=text, kind=kind, heading=heading))
        lines.append(text)

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text or "").strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        rendered = "\n".join(rows)
        if rendered:
            blocks.append(TextBlock(text=rendered, kind="table", heading=heading))
            lines.append(rendered)

    blob = _clean("\n".join(lines))
    if not blob:
        raise ValidationAppError("This Word document did not contain readable text.")

    pages = _split_pages(blob, blocks, heading)
    return pages


def extract_txt_pages(data: bytes) -> List[PageText]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")
    cleaned = _clean(text)
    if not cleaned:
        raise ValidationAppError("This text file is empty.")
    heading = None
    blocks: List[TextBlock] = []
    for raw in re.split(r"\n{2,}", cleaned):
        piece = raw.strip()
        if not piece:
            continue
        kind = "paragraph"
        first = piece.splitlines()[0].strip()
        if first.startswith("#") or (len(first) < 80 and first.isupper()):
            kind = "heading"
            heading = first.lstrip("# ").strip()[:180]
        elif re.match(r"^\s*([•\-\*]|\d+[\.)])\s+", piece):
            kind = "list"
        blocks.append(TextBlock(text=piece, kind=kind, heading=heading))
    return _split_pages(cleaned, blocks, heading)


def _split_pages(text: str, blocks: List[TextBlock], heading: str | None) -> List[PageText]:
    # TXT/DOCX have no native pages; keep ~1800 characters per synthetic page for citations.
    size = 1800
    pages: List[PageText] = []
    start = 0
    number = 1
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            break_at = text.rfind("\n\n", start, end)
            if break_at > start + size // 3:
                end = break_at
        piece = text[start:end].strip()
        if piece:
            page_blocks = [block for block in blocks if block.text and block.text in piece]
            pages.append(
                PageText(
                    page_number=number,
                    text=piece,
                    heading=heading,
                    blocks=page_blocks or [TextBlock(text=piece, kind="paragraph", heading=heading)],
                )
            )
            number += 1
        start = end if end > start else start + size
    return pages or [
        PageText(page_number=1, text=text, heading=heading, blocks=blocks)
    ]


def _clean(value: str) -> str:
    value = value.replace("\x00", " ").replace("\r\n", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()
